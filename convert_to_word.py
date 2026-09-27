import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

doc = docx.Document()

# Page Margins: 1 inch all around
for section in doc.sections:
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

# Colors
PURPLE = RGBColor(106, 27, 154)     # #6A1B9A
DARK_TEXT = RGBColor(33, 33, 33)    # #212121
GRAY_TEXT = RGBColor(117, 117, 117) # #757575
ACCENT = RGBColor(74, 20, 140)

# Set base font
style = doc.styles['Normal']
font = style.font
font.name = 'Calibri'
font.size = Pt(11)
font.color.rgb = DARK_TEXT

# Title
title_p = doc.add_paragraph()
title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
title_run = title_p.add_run("🎓 TCA APP — ALL DEFENSE QUESTIONS & ANSWERS")
title_run.font.size = Pt(20)
title_run.font.bold = True
title_run.font.color.rgb = PURPLE

# Subtitle
sub_p = doc.add_paragraph()
sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
sub_run = sub_p.add_run("Official Capstone / Thesis Defense & Project Pitching Reviewer\n(Pinaka-Simple, Direct to the Point, ug Dali Masag-ulo)")
sub_run.font.size = Pt(12)
sub_run.font.italic = True
sub_run.font.color.rgb = GRAY_TEXT

doc.add_paragraph() # Spacing

# Read markdown content
with open(r"c:\MOBILEAPP\TCA_APP\DEFENSE_ALL_QUESTIONS_AND_ANSWERS.md", "r", encoding="utf-8") as f:
    lines = f.readlines()

for line in lines:
    line = line.strip()
    if not line:
        continue
    if line.startswith("# ") or "STUDY GUIDE" in line:
        continue # handled above
    elif line.startswith("## "):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(line.replace("## ", ""))
        run.font.size = Pt(15)
        run.font.bold = True
        run.font.color.rgb = PURPLE
        # Add bottom border/separator
        pBorder = parse_xml(r'<w:pBdr {}><w:bottom w:val="single" w:sz="12" w:space="4" w:color="6A1B9A"/></w:pBdr>'.format(nsdecls('w')))
        p._p.get_or_add_pPr().append(pBorder)
    elif line.startswith("### "):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(2)
        q_text = line.replace("### ", "")
        run = p.add_run(q_text)
        run.font.size = Pt(12)
        run.font.bold = True
        run.font.color.rgb = ACCENT
    elif line.startswith("* **Tubag:**"):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.space_after = Pt(6)
        
        # Parse bold formatting in answer
        ans_text = line.replace("* **Tubag:**", "").strip()
        ans_lbl = p.add_run("👉 Tubag: ")
        ans_lbl.font.bold = True
        ans_lbl.font.color.rgb = PURPLE
        
        # split by ** for bold
        parts = ans_text.split("**")
        is_bold = False
        for part in parts:
            if part:
                r = p.add_run(part)
                r.font.size = Pt(11)
                r.font.bold = is_bold
                if is_bold:
                    r.font.color.rgb = DARK_TEXT
            is_bold = not is_bold
    elif line.startswith("* **") and ":**" in line:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.space_after = Pt(3)
        parts = line.split(":**")
        key = parts[0].replace("* **", "").strip()
        val = parts[1].strip() if len(parts) > 1 else ""
        
        r1 = p.add_run(f"• {key}: ")
        r1.font.bold = True
        r1.font.color.rgb = PURPLE
        
        val_parts = val.split("**")
        is_bold = False
        for vp in val_parts:
            if vp:
                r = p.add_run(vp)
                r.font.bold = is_bold
            is_bold = not is_bold
    elif line.startswith("---"):
        pass
    else:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.space_after = Pt(4)
        p.add_run(line)

output_path = r"c:\MOBILEAPP\TCA_APP\DEFENSE_ALL_QUESTIONS_AND_ANSWERS.docx"
doc.save(output_path)
print("Saved docx to:", output_path)
