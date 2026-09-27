import docx
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

file_path = r"C:\Users\antonette\OneDrive\Documents\NGANONG DGHN NGA ALGO ANG GI GAMIT NINYU.docx"

try:
    doc = docx.Document(file_path)
except Exception:
    doc = docx.Document()

PURPLE = RGBColor(106, 27, 154)     # #6A1B9A
DARK_TEXT = RGBColor(33, 33, 33)    # #212121
ACCENT = RGBColor(74, 20, 140)

# Add Page Break before adding the new comprehensive section
doc.add_page_break()

# Title for the newly added complete guide
p_title = doc.add_paragraph()
r_title = p_title.add_run("🎓 COMPLETE DEFENSE & PITCH REVIEWER GUIDE (Q&A 1 TO 40)")
r_title.font.size = Pt(16)
r_title.font.bold = True
r_title.font.color.rgb = PURPLE

# Read the comprehensive markdown content
with open(r"c:\MOBILEAPP\TCA_APP\DEFENSE_ALL_QUESTIONS_AND_ANSWERS.md", "r", encoding="utf-8") as f:
    lines = f.readlines()

for line in lines:
    line = line.strip()
    if not line:
        continue
    if line.startswith("# ") or "STUDY GUIDE" in line:
        continue
    elif line.startswith("## "):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(line.replace("## ", ""))
        run.font.size = Pt(14)
        run.font.bold = True
        run.font.color.rgb = PURPLE
        pBorder = parse_xml(r'<w:pBdr {}><w:bottom w:val="single" w:sz="12" w:space="4" w:color="6A1B9A"/></w:pBdr>'.format(nsdecls('w')))
        p._p.get_or_add_pPr().append(pBorder)
    elif line.startswith("### "):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(line.replace("### ", ""))
        run.font.size = Pt(11.5)
        run.font.bold = True
        run.font.color.rgb = ACCENT
    elif line.startswith("* **Tubag:**"):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.space_after = Pt(5)
        ans_text = line.replace("* **Tubag:**", "").strip()
        ans_lbl = p.add_run("👉 Tubag: ")
        ans_lbl.font.bold = True
        ans_lbl.font.color.rgb = PURPLE
        parts = ans_text.split("**")
        is_bold = False
        for part in parts:
            if part:
                r = p.add_run(part)
                r.font.bold = is_bold
            is_bold = not is_bold
    elif line.startswith("* **") and ":**" in line:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.space_after = Pt(3)
        parts = line.split(":**")
        key = parts[0].replace("* **", "").strip()
        val = parts[1].strip() if len(parts) > 1 else ""
        r1 = p.add_run(f"• {key}: ")
        r1.font.bold = True
        r1.font.color.rgb = PURPLE
        val_parts = val.split("**")
        is_bold = False
        for vp in val_parts:
            if vp:
                r = p.add_run(vp)
                r.font.bold = is_bold
            is_bold = not is_bold
    elif line.startswith("---"):
        pass
    else:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.space_after = Pt(4)
        p.add_run(line)

doc.save(file_path)
print("Successfully appended to:", file_path)
