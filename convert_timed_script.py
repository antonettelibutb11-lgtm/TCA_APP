import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

doc = docx.Document()

for section in doc.sections:
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

PURPLE = RGBColor(106, 27, 154)     # #6A1B9A
DARK_TEXT = RGBColor(33, 33, 33)    # #212121
GRAY_TEXT = RGBColor(117, 117, 117) # #757575
ACCENT = RGBColor(74, 20, 140)

style = doc.styles['Normal']
font = style.font
font.name = 'Calibri'
font.size = Pt(11)
font.color.rgb = DARK_TEXT

title_p = doc.add_paragraph()
title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
t_run = title_p.add_run("⏱️ TCA APP — 5-MINUTE EXACT TIMED DEMO SCRIPT")
t_run.font.size = Pt(18)
t_run.font.bold = True
t_run.font.color.rgb = PURPLE

sub_p = doc.add_paragraph()
sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
s_run = sub_p.add_run("Total Defense Time: 10 Minutes (5 Minutes Live Demo + 5 Minutes Q&A)\nDesigned for 2 Presenters (Speaker 1: Student View | Speaker 2: Admin View)")
s_run.font.size = Pt(11)
s_run.font.italic = True
s_run.font.color.rgb = GRAY_TEXT

doc.add_paragraph()

with open(r"c:\MOBILEAPP\TCA_APP\TIMED_5_MINUTE_DEMO_SCRIPT.md", "r", encoding="utf-8") as f:
    lines = f.readlines()

for line in lines:
    line = line.strip()
    if not line or line.startswith("# "):
        continue
    elif line.startswith("### "):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(line.replace("### ", ""))
        r.font.size = Pt(14)
        r.font.bold = True
        r.font.color.rgb = PURPLE
        pBorder = parse_xml(r'<w:pBdr {}><w:bottom w:val="single" w:sz="12" w:space="4" w:color="6A1B9A"/></w:pBdr>'.format(nsdecls('w')))
        p._p.get_or_add_pPr().append(pBorder)
    elif line.startswith("#### "):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(line.replace("#### ", ""))
        r.font.size = Pt(12)
        r.font.bold = True
        r.font.color.rgb = ACCENT
    elif line.startswith("> **"):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.2)
        p.paragraph_format.space_after = Pt(4)
        clean = line.replace("> ", "").strip()
        parts = clean.split("**")
        is_bold = False
        for part in parts:
            if part:
                r = p.add_run(part)
                r.font.bold = is_bold
                if is_bold:
                    r.font.color.rgb = PURPLE
            is_bold = not is_bold
    elif line.startswith("*("):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.2)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(line)
        r.font.italic = True
        r.font.color.rgb = GRAY_TEXT
    elif line.startswith("---"):
        pass
    else:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.2)
        p.paragraph_format.space_after = Pt(4)
        p.add_run(line)

out_file = r"c:\MOBILEAPP\TCA_APP\TIMED_5_MINUTE_DEMO_SCRIPT.docx"
doc.save(out_file)
print("Saved to:", out_file)
